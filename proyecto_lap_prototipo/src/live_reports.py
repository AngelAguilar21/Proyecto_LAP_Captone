"""Export observed session data with dependency-free PDF fallback."""
import csv
from datetime import datetime, timezone
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import textwrap


def report_data(config, state, kind):
    if not state.get("session"):
        raise ValueError("Inicia una sesión para generar un reporte.")
    if kind=="cameras":
        headers=['Cámara','Última muestra (s)','Personas en muestra','Máximo observado','Instante del máximo (s)','Promedio visible','Entradas','Salidas']
        names={c['id']:c.get('name',c['id']) for c in config['cameras']}
        rows=[]
        for cid,a in state.get('cameraAnalytics',{}).items():
            o=a.get('occupancy',{});lines=a.get('crossings',[])
            rows.append([names.get(cid,cid),a.get('t',state['t']),o.get('count'),o.get('peak'),o.get('peakAt'),round(o['mean'],2) if o.get('mean') is not None else None,sum(l['entries'] for l in lines),sum(l['exits'] for l in lines)])
        title='Resumen de monitoreo por cámara'
    elif kind in ("zones", "comparison"):
        headers=["Sector","Nivel","Referencia","Sector X","Sector Y","Personas x segundos","Pico observado","IDs observados","Prioridad"]
        cameras=[v for v in config['cameras'] if v.get('planId','custom')==state.get('planId',config.get('planId','custom'))]
        def reference(cell):
            if not cameras:return config.get('floor','Plano')
            near=min(cameras,key=lambda v:(v['x']-cell['x'])**2+(v['y']-cell['y'])**2)
            return 'Cerca de '+near.get('name',near['id'])
        rows=[[f'S{i+1}',state.get('planId',config.get('planId','custom')),reference(c),c["x"],c["y"],round(c["seconds"],3),c["peak"],c.get("visits",0),i+1] for i,c in enumerate(state["analytics"]["heat"])]
        title="Zonas de ocupación" if kind=="zones" else "Comparación de sectores"
    elif kind=="access-events":
        headers=['Cámara','Local / acceso','Movimiento','Instante de fuente (s)']
        names={c['id']:c.get('name',c['id']) for c in config['cameras']}
        rows=[[names.get(cid,cid),f"{l['place']['name']} / {l['name']}" if l.get('place') else l['name'],'Entrada' if e['direction']=='entries' else 'Salida',e['t']] for cid,a in state.get('cameraAnalytics',{}).items() for l in a.get('crossings',[]) for e in l.get('events',[])]
        title='Eventos de entrada y salida por local (hasta 1000 por acceso)'
    elif kind=="occupancy":
        headers=["Zona","Ocupación actual","Personas x segundos","Pico observado","IDs observados"]
        rows=[[z["name"],z["count"],round(z.get("seconds",0),3),z.get("peak",0),z.get("visits",0)] for z in state["analytics"]["zones"]]
        title="Ocupación por zona"
    elif kind=="crossings":
        headers=["Zona", "Entradas observadas", "Salidas observadas", "Último cruce (s de fuente)"]
        rows=[[z["name"],z["entries"],z["exits"],z["lastCrossing"]] for z in state["analytics"].get("flow",[])]
        for cid,a in state.get('cameraAnalytics',{}).items():
            rows.extend([[f"{cid} / {line['name']}",line['entries'],line['exits'],line['lastCrossing']] for line in a.get('crossings',[])])
        title="Entradas y salidas por zona y acceso"
    elif kind=="flow":
        bins={}
        for sample in state.get("series",[]):
            hour=int(sample["t"]//3600)
            item=bins.setdefault(hour,{"sum":0,"n":0,"peak":0})
            item["sum"]+=sample["count"]
            item["n"]+=1
            item["peak"]=max(item["peak"],sample["count"])
        headers=["Hora relativa de fuente","Personas observadas promedio","Pico observado","Muestras"]
        rows=[[h,round(v["sum"]/v["n"],3),v["peak"],v["n"]] for h,v in sorted(bins.items())]
        title="Flujo por hora relativa de fuente"
    elif kind=="trajectories":
        if state["status"] not in ("running","paused"):
            raise ValueError("Los recorridos individuales se eliminan al finalizar. Exporta durante una sesión activa o pausada.")
        headers=["ID temporal","Cámara actual","Asociación","Tiempo de fuente (s)","X","Y"]
        unique={p["id"]:p for p in state["people"]}
        rows=[[p["id"],p["camera"],p["association"],round(h[2],3),round(h[0],4),round(h[1],4)] for p in unique.values() for h in p["history"]]
        title="Trayectorias recientes de la sesión"
    else:
        raise ValueError("Tipo de reporte desconocido.")
    return {"title":title,"headers":headers,"rows":rows,"airport":(config.get("airport") or "Espacio sin nombre"),"session":state["session"],"mode":"SIMULACIÓN SINTÉTICA" if state.get("mode")=="demo" else "VIDEOS DE PRUEBA - PLANO ILUSTRATIVO" if state.get("testRun") else "FUENTE REAL - SIN VALIDACIÓN DE PRECISIÓN", "unit":"metros" if config["unit"]=="meters" else "unidades relativas","seconds":round(state["t"],2),"generated":datetime.now(timezone.utc).isoformat(),"note":"Datos observados; asociaciones estimadas. Ocupación no equivale a rentabilidad. Flujo horario calculado sobre las muestras retenidas (hasta 3600)."}


def _median(values):
    values=[v for v in values if v is not None]
    if not values:
        return 0
    ordered=sorted(values)
    n=len(ordered)
    mid=n//2
    return ordered[mid] if n%2 else (ordered[mid-1]+ordered[mid])/2


def _verdict(zone, median_seconds, median_visits):
    if zone.get("alert"):
        return "Aglomeración recurrente, revisar flujo"
    high_traffic=(zone.get("visits") or 0)>=median_visits
    high_dwell=(zone.get("seconds") or 0)>=median_seconds
    if high_traffic and high_dwell:
        return "Mayor tráfico y permanencia"
    if high_traffic:
        return "Zona de paso, poca permanencia"
    if high_dwell:
        return "Punto de permanencia, poco tráfico"
    return "Tráfico moderado"


def business_report_data(config, state):
    if not state.get("session"):
        raise ValueError("Inicia una sesión para generar un reporte.")
    zones=state["analytics"]["zones"]
    median_seconds=_median([z.get("seconds",0) for z in zones])
    median_visits=_median([z.get("visits",0) for z in zones])
    ranking=sorted(zones,key=lambda z:-(z.get("seconds") or 0))
    top=ranking[0] if ranking else None
    series=state.get("series",[])
    peak=max(series,key=lambda s:s["count"]) if series else None
    names={c['id']:c.get('name',c['id']) for c in config['cameras']}
    crossings=[{**l,"camera":names.get(cid,cid)} for cid,a in state.get('cameraAnalytics',{}).items() for l in a.get('crossings',[])]
    bins={}
    for sample in series:
        hour=int(sample["t"]//3600)
        item=bins.setdefault(hour,{"sum":0,"n":0,"peak":0})
        item["sum"]+=sample["count"];item["n"]+=1;item["peak"]=max(item["peak"],sample["count"])
    hourly=[(h,round(v["sum"]/v["n"],2),v["peak"],v["n"]) for h,v in sorted(bins.items())]
    recommendations=[]
    if top and (top.get("seconds") or 0)>0:
        recommendations.append(f"{top['name']} concentra la mayor permanencia acumulada de la sesión ({round(top.get('seconds',0))} persona-segundo). Es la primera candidata para evaluar un nuevo punto comercial o ampliar el actual, contrastando con otras temporadas y horarios antes de decidir.")
    alert_zones=[z for z in zones if z.get("alert")]
    if alert_zones:
        recommendations.append(f"{', '.join(z['name'] for z in alert_zones)} registra aglomeraciones recurrentes. Antes de invertir en espacio conviene revisar señalización o flujo: una zona congestionada no siempre es rentable.")
    pass_zones=[z for z in zones if (z.get("visits") or 0)>=median_visits and (z.get("seconds") or 0)<median_seconds and z is not top]
    if pass_zones:
        recommendations.append(f"{', '.join(z['name'] for z in pass_zones)} tiene tráfico alto pero poca permanencia. Encaja mejor con publicidad o señalización de paso que con retail que dependa de tiempo de exposición.")
    active_cameras=len([c for c in state.get('cameras',[]) if c.get('status')=='live']) if state['status'] in ('running','paused') else len(state.get('cameraAnalytics',{}))
    total_cameras=len(config['cameras'])
    if total_cameras and active_cameras<total_cameras:
        recommendations.append(f"Solo {active_cameras} de {total_cameras} cámaras aportaron resultados en esta sesión. Antes de comparar zonas con confianza para una decisión comercial conviene ampliar la cobertura o repetir la medición con todas las fuentes activas.")
    if not recommendations:
        recommendations.append("Aún no hay suficientes datos en esta sesión para generar recomendaciones. Estas aparecen a medida que se observan zonas y accesos.")
    peak_time=f"a las {int(peak['t']//60):02d}:{int(peak['t']%60):02d}" if peak else "sin muestras aún"
    return {
        "report_kind":"business",
        "title":"Reporte comercial",
        "airport":(config.get("airport") or "Espacio sin nombre"),
        "session":state["session"],
        "mode":"SIMULACIÓN SINTÉTICA" if state.get("mode")=="demo" else "VIDEOS DE PRUEBA - PLANO ILUSTRATIVO" if state.get("testRun") else "FUENTE REAL - SIN VALIDACIÓN DE PRECISIÓN",
        "unit":"metros" if config["unit"]=="meters" else "unidades relativas",
        "seconds":round(state["t"],2),
        "generated":datetime.now(timezone.utc).isoformat(),
        "kpis":[
            ("Ocupación pico", str(peak["count"]) if peak else "—", peak_time),
            ("Zona con más tráfico", top["name"] if top else "—", f"{round(top.get('seconds',0))} persona-seg." if top else "define zonas en el plano"),
            ("Permanencia media", f"{(state.get('totals') or {}).get('meanObservedSeconds',0):.1f} s", "por persona observada"),
            ("Aglomeraciones", str((state.get('totals') or {}).get('alerts',0)), "episodios en la sesión"),
        ],
        "ranking":[(z["name"], z.get("visits") if z.get("visits") is not None else "—",
                    f"{(z['seconds']/z['visits']):.0f} s" if z.get("visits") else "—",
                    z.get("peak") if z.get("peak") is not None else "—",
                    _verdict(z,median_seconds,median_visits)) for z in ranking],
        "crossings":[(f"{l['camera']} / {l.get('name','')}", l.get('entries',0), l.get('exits',0)) for l in crossings],
        "hourly":hourly,
        "recommendations":recommendations,
        "note":"Estimaciones de una sesión de tracking anónimo, sin identificación personal. No sustituyen un estudio de rentabilidad; son un punto de partida para decidir dónde mirar primero.",
    }


def business_pdf_bytes(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from xml.sax.saxutils import escape
    surface=colors.HexColor("#F0F5F8");line=colors.HexColor("#C7D9E3")
    text=colors.HexColor("#173A50");muted=colors.HexColor("#5B7386");amber=colors.HexColor("#0A3958");blue=colors.HexColor("#0C4B74");white=colors.white
    buffer=io.BytesIO()
    document=SimpleDocTemplate(buffer,pagesize=A4,leftMargin=32,rightMargin=32,topMargin=30,bottomMargin=36,title=data["title"],author="AeroTrack")
    title_style=ParagraphStyle('t',fontName='Helvetica-Bold',fontSize=22,textColor=amber,leading=26)
    sub_style=ParagraphStyle('s',fontName='Helvetica',fontSize=10,textColor=muted,leading=14)
    h2_style=ParagraphStyle('h2',fontName='Helvetica-Bold',fontSize=13,textColor=amber,spaceBefore=16,spaceAfter=7)
    body_style=ParagraphStyle('b',fontName='Helvetica',fontSize=9.5,textColor=text,leading=14)
    note_style=ParagraphStyle('n',fontName='Helvetica-Oblique',fontSize=8.5,textColor=muted,leading=12,spaceBefore=10)
    empty_style=ParagraphStyle('e',fontName='Helvetica-Oblique',fontSize=9.5,textColor=muted,leading=13)

    def paint(canvas,doc):
        canvas.saveState()
        canvas.setFont("Helvetica",8)
        canvas.setFillColor(muted)
        canvas.drawString(32,16,"AeroTrack · Reporte comercial · Sin identificación personal")
        canvas.drawRightString(A4[0]-32,16,f"Página {doc.page}")
        canvas.restoreState()

    def table_style(header_bg,header_color,rows):
        style=[('BACKGROUND',(0,0),(-1,0),header_bg),('TEXTCOLOR',(0,0),(-1,0),header_color),
               ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8.5),
               ('TEXTCOLOR',(0,1),(-1,-1),text),('ROWBACKGROUNDS',(0,1),(-1,-1),[white,surface]),
               ('GRID',(0,0),(-1,-1),0.5,line),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
               ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),
               ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8)]
        return TableStyle(style)

    flow=[Paragraph("AeroTrack | "+escape(data["title"]),title_style),
          Paragraph(escape(data["airport"])+" · "+escape(data["mode"]),sub_style),
          Spacer(1,4),
          Paragraph(f"Sesión {escape(data['session'])} · {data['seconds']} segundos de fuente · {escape(data['unit'])} · generado {escape(data['generated'])}",sub_style),
          Spacer(1,16)]

    kpi_rows=[[k[0].upper() for k in data["kpis"]],[k[1] for k in data["kpis"]],[k[2] for k in data["kpis"]]]
    kpi_table=Table(kpi_rows,colWidths=[document.width/4]*4)
    kpi_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),surface),
        ('TEXTCOLOR',(0,0),(-1,0),muted),('FONTSIZE',(0,0),(-1,0),7.5),
        ('TEXTCOLOR',(0,1),(-1,1),amber),('FONTSIZE',(0,1),(-1,1),18),('FONTNAME',(0,1),(-1,1),'Helvetica-Bold'),
        ('TEXTCOLOR',(0,2),(-1,2),muted),('FONTSIZE',(0,2),(-1,2),7),
        ('GRID',(0,0),(-1,-1),0.6,line),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
        ('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10)]))
    flow.append(kpi_table)

    flow.append(Paragraph("Zonas por oportunidad comercial",h2_style))
    if not data["ranking"]:
        flow.append(Paragraph("Define zonas en el plano para obtener este ranking.",empty_style))
    else:
        rows=[["Zona","Visitas","Permanencia media","Pico","Lectura"]]+[[str(v) for v in r] for r in data["ranking"]]
        t=Table(rows,colWidths=[document.width*0.24,document.width*0.13,document.width*0.19,document.width*0.13,document.width*0.31],repeatRows=1)
        t.setStyle(table_style(amber,white,rows))
        flow.append(t)

    flow.append(Paragraph("Tráfico por acceso",h2_style))
    if not data["crossings"]:
        flow.append(Paragraph("No hay accesos con cruces confirmados en esta sesión.",empty_style))
    else:
        rows=[["Cámara / acceso","Entradas","Salidas"]]+[[str(v) for v in r] for r in data["crossings"]]
        t=Table(rows,colWidths=[document.width*0.6,document.width*0.2,document.width*0.2],repeatRows=1)
        t.setStyle(table_style(blue,white,rows))
        flow.append(t)

    flow.append(Paragraph("Evolución por hora relativa de fuente",h2_style))
    if not data["hourly"]:
        flow.append(Paragraph("El gráfico aparecerá al acumular muestras durante la sesión.",empty_style))
    else:
        rows=[["Hora relativa","Promedio","Pico","Muestras"]]+[[str(v) for v in r] for r in data["hourly"]]
        t=Table(rows,colWidths=[document.width*0.25]*4,repeatRows=1)
        t.setStyle(table_style(blue,white,rows))
        flow.append(t)

    flow.append(Paragraph("Recomendaciones para LAP",h2_style))
    for r in data["recommendations"]:
        flow.append(Paragraph("• "+escape(r),body_style))
        flow.append(Spacer(1,5))
    flow.append(Paragraph(escape(data["note"]),note_style))

    document.build(flow,onFirstPage=paint,onLaterPages=paint)
    return buffer.getvalue()


def business_xlsx_bytes(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    alt=PatternFill("solid",fgColor="EDF4F8")
    navy=Font(name="Calibri",bold=True,color="0A3958");dark=Font(name="Calibri",color="173A50")
    muted=Font(name="Calibri",color="5B7386",italic=True)
    wb=Workbook();ws=wb.active;ws.title="Reporte comercial"
    row=1
    def title(text_,font):
        nonlocal row
        ws.cell(row,1,text_).font=font
        row+=1
    title("AeroTrack | "+data["title"],Font(name="Calibri",size=18,bold=True,color="0A3958"))
    title(data["mode"]+" · "+data["airport"],muted)
    title(f"Sesión {data['session']} · {data['seconds']} s · {data['unit']} · generado {data['generated']}",muted)
    row+=1
    title("Indicadores clave",navy)
    for label,value,note in data["kpis"]:
        ws.cell(row,1,label).font=dark;ws.cell(row,2,value).font=navy;ws.cell(row,3,note).font=muted
        row+=1
    row+=1
    title("Zonas por oportunidad comercial",navy)
    headers=["Zona","Visitas","Permanencia media","Pico","Lectura"]
    for i,h in enumerate(headers,1):
        c=ws.cell(row,i,h);c.font=Font(name="Calibri",bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="0C4B74")
    row+=1
    for r in data["ranking"]:
        for i,v in enumerate(r,1):
            c=ws.cell(row,i,v);c.font=dark
            if row%2==0:c.fill=alt
        row+=1
    row+=1
    title("Tráfico por acceso",navy)
    for i,h in enumerate(["Cámara / acceso","Entradas","Salidas"],1):
        c=ws.cell(row,i,h);c.font=Font(name="Calibri",bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="0C4B74")
    row+=1
    for r in data["crossings"]:
        for i,v in enumerate(r,1):
            c=ws.cell(row,i,v);c.font=dark
            if row%2==0:c.fill=alt
        row+=1
    row+=1
    title("Recomendaciones para LAP",navy)
    for r in data["recommendations"]:
        ws.cell(row,1,"• "+r).font=dark
        row+=1
    row+=1
    ws.cell(row,1,data["note"]).font=muted
    for col in "ABCDE":
        ws.column_dimensions[col].width=30
    output=io.BytesIO();wb.save(output)
    return output.getvalue()


def business_csv_bytes(data):
    output=io.StringIO(newline="")
    writer=csv.writer(output)
    writer.writerow(["AeroTrack",data["title"]])
    writer.writerow(["Origen",data["mode"],"Sesión",data["session"],"Unidad",data["unit"]])
    writer.writerow([])
    writer.writerow(["Indicadores clave"])
    for label,value,note in data["kpis"]:
        writer.writerow([label,value,note])
    writer.writerow([])
    writer.writerow(["Zonas por oportunidad comercial"])
    writer.writerow(["Zona","Visitas","Permanencia media","Pico","Lectura"])
    writer.writerows(data["ranking"])
    writer.writerow([])
    writer.writerow(["Tráfico por acceso"])
    writer.writerow(["Cámara / acceso","Entradas","Salidas"])
    writer.writerows(data["crossings"])
    writer.writerow([])
    writer.writerow(["Recomendaciones para LAP"])
    for r in data["recommendations"]:
        writer.writerow([r])
    writer.writerow([])
    writer.writerow([data["note"]])
    return output.getvalue().encode("utf-8-sig")


def safe_text(value):
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value


def csv_bytes(data):
    output=io.StringIO(newline="")
    writer=csv.writer(output)
    writer.writerow(["AeroTrack",data["title"]])
    writer.writerow(["Origen",data["mode"],"Sesión",data["session"],"Unidad",data["unit"]])
    writer.writerow(data["headers"])
    writer.writerows([[safe_text(v) for v in row] for row in data["rows"]])
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook()
    ws=wb.active
    ws.title="Reporte"
    columns=len(data["headers"])
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=columns)
    ws.cell(1,1,"AeroTrack | "+data["title"]).font=Font(name="Calibri",size=18,bold=True,color="0A3958")
    ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=columns)
    ws.cell(2,1,data["mode"]+" · "+data["airport"]).font=Font(size=10,color="456577")
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=columns)
    ws.cell(3,1,f"Sesión {data['session']} · {data['seconds']} segundos de fuente · {data['unit']}")
    ws.append([])
    for i,title in enumerate(data["headers"],1):
        cell=ws.cell(5,i,title)
        cell.fill=PatternFill("solid",fgColor="0C4B74")
        cell.font=Font(name="Calibri",bold=True,color="FFFFFF")
        cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[get_column_letter(i)].width=max(18,min(32,len(title)+3))
    ws.row_dimensions[5].height=32
    for ri,row in enumerate(data["rows"],6):
        for ci,value in enumerate(row,1):
            cell=ws.cell(ri,ci,safe_text(value))
            cell.font=Font(name="Calibri",size=11,color="173A50")
            if isinstance(value,float):
                cell.number_format="0.000"
            if ri%2==0:
                cell.fill=PatternFill("solid",fgColor="EDF4F8")
    ws.freeze_panes="A6"
    ws.auto_filter.ref=f"A5:{get_column_letter(columns)}{max(5,5+len(data['rows']))}"
    note=ws.cell(len(data["rows"])+8,1,data["note"])
    note.font=Font(size=10,color="526A79")
    ws.merge_cells(start_row=note.row,start_column=1,end_row=note.row,end_column=columns)
    note.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[note.row].height=38
    output=io.BytesIO()
    wb.save(output)
    return output.getvalue()


def pdf_bytes(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape,A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from xml.sax.saxutils import escape
    buffer=io.BytesIO()
    document=SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=34,rightMargin=34,topMargin=30,bottomMargin=30,title=data["title"],author="AeroTrack")
    styles=getSampleStyleSheet()
    styles["Title"].textColor=colors.HexColor("#0a3958")
    styles["Title"].fontSize=20
    styles["Normal"].fontSize=9
    flow=[Paragraph("AeroTrack | "+escape(data["title"]),styles["Title"]),Paragraph(escape(data["airport"])+" · "+escape(data["mode"]),styles["Normal"]),Spacer(1,8),Paragraph(f"Sesión {escape(data['session'])} · {data['seconds']} segundos de fuente · {escape(data['unit'])}",styles["Normal"]),Spacer(1,18)]
    rows=[data["headers"]]+data["rows"]
    formatted=[[Paragraph(escape(str(v)),styles["Normal"]) for v in row] for row in rows]
    if len(rows)==1:
        flow.append(Paragraph("Sin observaciones para este reporte.",styles["Normal"]))
    else:
        table=Table(formatted,colWidths=[document.width/len(data["headers"])]*len(data["headers"]),repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DBEAF3")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F5F8")]),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),("LINEBELOW",(0,0),(-1,0),.6,colors.HexColor("#2679AE"))]))
        flow.append(table)
    flow.extend([Spacer(1,16),Paragraph(escape(data["note"]),styles["Normal"])])
    def footer(canvas,doc):
        canvas.setFont("Helvetica",8)
        canvas.setFillColor(colors.HexColor("#608092"))
        canvas.drawString(34,16,"AeroTrack · Información de sesión · Sin identificación personal")
        canvas.drawRightString(landscape(A4)[0]-34,16,f"Página {doc.page}")
    document.build(flow,onFirstPage=footer,onLaterPages=footer)
    return buffer.getvalue()


def simple_pdf_bytes(data):
    """Genera un PDF textual válido cuando ReportLab no está instalado."""
    lines = [f"AeroTrack | {data['title']}", f"{data['airport']} | {data['mode']}",
             f"Sesion {data['session']} | {data['seconds']} segundos | {data['unit']}", ""]
    if data.get("report_kind") == "business":
        lines.append("INDICADORES")
        lines.extend(f"{label}: {value} | {detail}" for label, value, detail in data.get("kpis", []))
        for title, key in (("RANKING", "ranking"), ("ACCESOS", "crossings"), ("FRANJAS HORARIAS", "hourly")):
            lines.extend(("", title))
            lines.extend(" | ".join(map(str, row)) for row in data.get(key, []))
        lines.extend(("", "RECOMENDACIONES"))
        lines.extend(data.get("recommendations", []))
    else:
        lines.append(" | ".join(map(str, data.get("headers", []))))
        lines.extend(" | ".join(map(str, row)) for row in data.get("rows", []))
    lines.extend(("", str(data.get("note", ""))))
    wrapped = []
    for line in lines:
        wrapped.extend(textwrap.wrap(str(line), width=118, replace_whitespace=True,
                                     drop_whitespace=True) or [""])
    pages = [wrapped[i:i + 42] for i in range(0, len(wrapped), 42)] or [[""]]

    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    page_ids = []
    for index, page in enumerate(pages):
        page_id, content_id = 4 + index * 2, 5 + index * 2
        page_ids.append(page_id)
        content = bytearray(b"BT\n/F1 9 Tf\n34 558 Td\n12 TL\n")
        for line in page:
            literal = line.encode("cp1252", "replace").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
            content.extend(b"(" + literal + b") Tj\nT*\n")
        content.extend(b"ET\n")
        objects[page_id] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] "
                            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>").encode()
        objects[content_id] = b"<< /Length %d >>\nstream\n" % len(content) + content + b"endstream"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0] * (max(objects) + 1)
    for object_id in range(1, len(offsets)):
        offsets[object_id] = len(output)
        output.extend(f"{object_id} 0 obj\n".encode() + objects[object_id] + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def export(config,state,kind,format):
    if kind=="business":
        data=business_report_data(config,state)
        csv_fn,xlsx_fn,pdf_fn=business_csv_bytes,business_xlsx_bytes,business_pdf_bytes
    else:
        data=report_data(config,state,kind)
        csv_fn,xlsx_fn,pdf_fn=csv_bytes,xlsx_bytes,pdf_bytes
    if format=="csv":
        return csv_fn(data),"text/csv; charset=utf-8"
    function={"xlsx":xlsx_fn,"pdf":pdf_fn}.get(format)
    if not function:
        raise ValueError("Formato desconocido.")
    try:
        result=function(data)
    except ImportError:
        if format == "pdf":
            return simple_pdf_bytes(data), "application/pdf"
        runtime=os.environ.get("AEROTRACK_DOCUMENT_PYTHON")
        if not runtime or not Path(runtime).is_file():
            raise ValueError("Este reporte requiere reportlab (PDF) u openpyxl (Excel), o un AEROTRACK_DOCUMENT_PYTHON con esas dependencias.")
        process=subprocess.run([runtime,str(Path(__file__).resolve()),format],input=json.dumps(data).encode(),capture_output=True,timeout=45,creationflags=getattr(subprocess,"CREATE_NO_WINDOW",0))
        if process.returncode:
            raise ValueError("No se pudo generar el documento. Comprueba las dependencias de exportación.")
        result=process.stdout
    return result,"application/pdf" if format=="pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


if __name__=="__main__":
    data=json.loads(sys.stdin.buffer.read())
    business=data.get("report_kind")=="business"
    pdf_fn=business_pdf_bytes if business else pdf_bytes
    xlsx_fn=business_xlsx_bytes if business else xlsx_bytes
    sys.stdout.buffer.write((pdf_fn if sys.argv[1]=="pdf" else xlsx_fn)(data))
