"""Repeatable commercial rehearsal. Uploaded history never overwrites operational sales."""
import base64
import csv
import hashlib
import io
import json
import math
import statistics
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

from commercial import LIMA

COLUMNS = ['negocio_id', 'fecha', 'hora', 'ventas_pen', 'transacciones', 'entradas', 'cobertura_min']


def setup(con):
    con.execute('''CREATE TABLE IF NOT EXISTS commercial_trials (
      id TEXT PRIMARY KEY, negocio_id TEXT NOT NULL, session TEXT NOT NULL,
      created TEXT NOT NULL, result TEXT NOT NULL)''')
    con.commit()


def read_history(encoded, filename, business_id):
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError('No se pudo leer el archivo. Selecciónalo de nuevo.') from exc
    if not raw or len(raw) > 2_000_000:
        raise ValueError('Usa un archivo de hasta 2 MB.')
    try:
        if filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 8_000_000 or len(archive.infolist()) > 200:
                    raise ValueError('El Excel contiene demasiados datos. Usa la plantilla.')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if 'Historico' not in book.sheetnames:
                    raise ValueError('El Excel debe tener una hoja Historico con los encabezados de la plantilla en la primera fila.')
                sheet = book['Historico']
                if sheet.max_row is None or sheet.max_column is None:
                    sheet.calculate_dimension(force=True)
                if sheet.max_row > 1001 or sheet.max_column > 20:
                    raise ValueError('La hoja Historico admite hasta 1.000 filas y 20 columnas.')
                cells = list(sheet.iter_rows(values_only=True))
                fields = [str(v or '').strip() for v in cells[0]] if cells else []
                records = [dict(zip(fields, row)) for row in cells[1:] if any(v is not None for v in row)]
            finally:
                book.close()
        elif filename.lower().endswith('.csv'):
            text = raw.decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(text), delimiter=';' if ';' in text.splitlines()[0] else ',')
            fields = reader.fieldnames or []
            records = list(reader)
        else:
            raise ValueError('Selecciona un Excel .xlsx o un CSV UTF-8.')
    except (zipfile.BadZipFile, UnicodeError, KeyError, IndexError) as exc:
        raise ValueError('Archivo inválido. Usa el ejemplo de la tienda seleccionada.') from exc
    if len(set(fields)) != len(fields) or not set(COLUMNS).issubset(fields):
        raise ValueError('Columnas requeridas: ' + ', '.join(COLUMNS))
    if not 1 <= len(records) <= 1000:
        raise ValueError('El histórico debe contener entre 1 y 1.000 filas.')
    rows, seen = [], set()
    for number, record in enumerate(records, 2):
        try:
            if str(record['negocio_id']).strip() != business_id:
                raise ValueError('el archivo pertenece a otra tienda; selecciona la tienda correcta')
            day = record['fecha']
            if isinstance(day, datetime):
                day = day.date().isoformat()
            if datetime.strptime(str(day), '%Y-%m-%d').date().isoformat() != day:
                raise ValueError('fecha debe ser AAAA-MM-DD')
            values = {}
            for name in COLUMNS[2:]:
                value = record[name]
                if isinstance(value, bool):
                    raise ValueError(name + ' debe ser numérico')
                value = float(str(value).replace(',', '.'))
                if not math.isfinite(value) or value < 0:
                    raise ValueError(name + ' debe ser finito y no negativo')
                values[name] = value
            if any(values[k] != int(values[k]) for k in ('hora', 'entradas', 'transacciones')):
                raise ValueError('hora, entradas y transacciones deben ser enteros')
            if values['hora'] > 23 or values['cobertura_min'] > 60:
                raise ValueError('hora debe ser 0–23 y cobertura_min 0–60')
            key = (day, int(values['hora']))
            if key in seen:
                raise ValueError('fecha y hora duplicadas')
            seen.add(key)
            rows.append(dict(negocio_id=business_id, fecha=day, **values))
        except (ValueError, TypeError, KeyError) as exc:
            raise ValueError(f'Fila {number}: {exc}.') from exc
    return rows, hashlib.sha256(raw).hexdigest()


def example_csv(business_id, day, hour):
    """Two periods, eight previous comparable dates. Explicitly fictitious example."""
    origin = datetime.fromisoformat(day) + timedelta(hours=hour)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(COLUMNS)
    for offset in (0, 1):
        for week in range(8, 0, -1):
            date = origin + timedelta(hours=offset) - timedelta(weeks=week)
            entries = 40 + 3 * week + 6 * offset
            transactions = round(entries * .2)
            writer.writerow([business_id, date.date().isoformat(), date.hour, transactions * 25, transactions, entries, 60])
    return output.getvalue()


def observation(meta, business):
    origin = datetime.fromisoformat(meta.get('config', {}).get('recordingStartedAt') or meta['created']).astimezone(LIMA)
    remaining = (origin.replace(minute=0, second=0, microsecond=0) + timedelta(hours=1) - origin).total_seconds()
    cameras = {a['cameraId'] for a in business['accesses']}
    times = [float(meta.get('cameraAnalytics', {}).get(cid, {}).get('t', 0)) for cid in cameras]
    coverage = max(0, min([float(meta.get('end', 0)), remaining, *times])) if times else 0
    # Analytics use 'time' in some legacy sessions.
    if coverage == 0 and cameras:
        times = [float(meta.get('cameraAnalytics', {}).get(cid, {}).get('time', 0)) for cid in cameras]
        coverage = max(0, min([float(meta.get('end', 0)), remaining, *times]))
    events = [e for e in business['events'] if 0 <= float(e['t']) <= coverage and float(e['t']) < remaining]
    return dict(session=meta['session'], date=origin.date().isoformat(), hour=origin.hour,
                coverage=coverage, entries=sum(e['direction'] == 'entries' for e in events),
                exits=sum(e['direction'] == 'exits' for e in events), cameras=sorted(cameras))


def calculate(rows, observed):
    def comparable(day, hour):
        weekday = datetime.fromisoformat(day).weekday()
        return [r for r in rows if r['fecha'] < day and int(r['hora']) == hour
                and datetime.fromisoformat(r['fecha']).weekday() == weekday
                and r['cobertura_min'] >= 55 and r['entradas'] > 0]
    current = comparable(observed['date'], observed['hour'])
    target = datetime.fromisoformat(observed['date']) + timedelta(hours=observed['hour'] + 1)
    future = comparable(target.date().isoformat(), target.hour)
    rate = statistics.median(r['ventas_pen'] / r['entradas'] for r in current) if len(current) >= 3 else None
    next_rate = statistics.median(r['ventas_pen'] / r['entradas'] for r in future) if len(future) >= 3 else None
    expected = round(statistics.median(r['entradas'] for r in future)) if next_rate is not None else None
    return dict(**observed, days=len(current), nextDays=len(future), rate=rate,
                estimate=round(observed['entries'] * rate, 2) if rate is not None and observed['coverage'] > 0 else None,
                nextTarget=target.isoformat(), expectedEntries=expected,
                nextRate=next_rate, forecast=round(expected * next_rate, 2) if next_rate is not None else None,
                historicalConversion=round(100 * sum(r['transacciones'] for r in current) / sum(r['entradas'] for r in current), 1) if len(current) >= 3 else None,
                historyTransactions=sum(r['transacciones'] for r in current), historyEntries=sum(r['entradas'] for r in current),
                conversion=None, history=rows, origin='Ensayo con archivo; no modifica ventas registradas')


def save(con, result, business, filename, fingerprint):
    setup(con)
    result = dict(result, id=uuid.uuid4().hex[:12], businessId=business['id'], businessName=business['nombre'],
                  filename=filename[:120], fingerprint=fingerprint, created=datetime.now(timezone.utc).isoformat())
    with con:
        con.execute('INSERT INTO commercial_trials VALUES (?,?,?,?,?)',
                    (result['id'], business['id'], result['session'], result['created'], json.dumps(result, ensure_ascii=False)))
    return result


def saved(con):
    setup(con)
    return [json.loads(row[0]) for row in con.execute('SELECT result FROM commercial_trials ORDER BY created DESC LIMIT 20')]


def session_description(meta):
    cameras = meta.get('cameras', [])
    kinds = {c.get('sourceKind') or ('live' if isinstance(c.get('source'), int) or str(c.get('source', '')).startswith(('http://', 'https://', 'rtsp://', 'rtmp://')) else 'recording') for c in cameras}
    return dict(cameraNames=[c.get('name') or c['id'] for c in cameras],
                duration=max(0, float(meta.get('end', 0))),
                sourceLabel='En vivo' if kinds == {'live'} else 'Videos grabados' if kinds == {'recording'} else 'Fuentes combinadas',
                statusLabel='Detenido por el operador' if meta.get('status') == 'stopped' else 'Análisis terminado')


def session_options(root, project_id, businesses, config):
    import business_catalog
    options = []
    for path in (root / 'data/replays').glob('*/manifest.json'):
        try:
            meta = json.loads(path.read_text(encoding='utf-8'))
            if meta.get('projectId') != project_id or meta.get('status') not in ('ended', 'stopped'):
                continue
            if not meta.get('config', {}).get('testRun'):
                continue
            traffic = business_catalog.traffic(businesses, meta.get('cameraAnalytics', {}),
                                               {**meta['config'], 'cameras': meta['cameras']}, config)
            measured = [dict(id=b['id'], name=b['nombre'], **observation(meta, b))
                        for b in traffic if b['entries'] is not None]
            measured = [b for b in measured if b['coverage'] > 0]
            if measured:
                options.append(dict(id=meta['session'], created=meta['created'], businesses=measured, **session_description(meta)))
        except (ValueError, OSError, KeyError, TypeError):
            continue
    return sorted(options, key=lambda s: s['created'], reverse=True)[:50]
